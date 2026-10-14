#!/usr/bin/env python3
"""
Test Scenario to CSV Converter Agent
Converts test scenarios from Excel/table format into Jira-ready CSV format
"""

import argparse
import csv
import sys
import re
from pathlib import Path
from typing import Dict, List, Optional

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False


class TestScenarioConverter:
    """Converts test scenarios to Jira-ready CSV format"""
    
    # CSV column headers (21 columns)
    HEADERS = [
        "Parent ID",
        "IssueType",
        "Assignee",
        "Fix Version",
        "User Story ID",
        "Scenario Type",
        "Region Tested",
        "TestType",
        "Test Scenario/Test Case ID",
        "Summary",
        "Acceptance Criteria",
        "Pre-Condition",
        "Test Scenario Description",
        "Test Steps",
        "Expected Result",
        "Test Data",
        "QA Planned Start Date",
        "QA Planned End Date",
        "QA Actual Start Date",
        "QA Actual End Date",
        "Priority"
    ]
    
    # Default values
    DEFAULTS = {
        "IssueType": "Test Case",
        "Scenario Type": "Functional",
        "Region Tested": "QA",
        "Priority": "Medium"
    }
    
    def __init__(self):
        self.test_cases = []
    
    def read_input(self, input_path: str) -> List[Dict]:
        """Read input from Excel, CSV, or text"""
        input_path = str(input_path).strip()
        
        # Check if it's a file path
        if Path(input_path).exists():
            if input_path.lower().endswith('.xlsx'):
                return self._read_excel(input_path)
            elif input_path.lower().endswith('.csv'):
                return self._read_csv(input_path)
            else:
                return self._read_text_file(input_path)
        else:
            # Treat as pasted text
            return self._parse_text(input_path)
    
    def _read_excel(self, filepath: str) -> List[Dict]:
        """Read data from Excel file"""
        if not HAS_OPENPYXL and not HAS_PANDAS:
            print("Error: openpyxl or pandas required for Excel support")
            print("Install with: pip install openpyxl pandas")
            sys.exit(1)
        
        try:
            if HAS_PANDAS:
                df = pd.read_excel(filepath, sheet_name=0)
                return df.fillna('').to_dict('records')
            else:
                wb = openpyxl.load_workbook(filepath)
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                data = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    data.append(dict(zip(headers, row)))
                return [d for d in data if any(d.values())]
        except Exception as e:
            print(f"Error reading Excel file: {e}")
            sys.exit(1)
    
    def _read_csv(self, filepath: str) -> List[Dict]:
        """Read data from CSV file"""
        try:
            data = []
            with open(filepath, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    data.append({k: (v or '') for k, v in row.items()})
            return data
        except Exception as e:
            print(f"Error reading CSV file: {e}")
            sys.exit(1)
    
    def _read_text_file(self, filepath: str) -> List[Dict]:
        """Read text file"""
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                text = f.read()
            return self._parse_text(text)
        except Exception as e:
            print(f"Error reading file: {e}")
            sys.exit(1)
    
    def _parse_text(self, text: str) -> List[Dict]:
        """Parse text/table format into test cases"""
        lines = text.strip().split('\n')
        data = []
        current_test = {}
        current_section = None
        
        for line in lines:
            line = line.strip()
            
            if not line:
                continue
            
            # Detect test case ID
            if any(line.upper().startswith(x) for x in ['TEST CASE', 'TC-', 'TEST-', '001', '002']):
                if current_test:
                    data.append(current_test)
                current_test = {'Test Case ID': line}
                current_section = None
            
            # Detect sections
            elif any(x in line.upper() for x in ['SCENARIO', 'PRE-CONDITION', 'TEST STEP', 'EXPECTED', 'TEST DATA']):
                current_section = line
                if 'Description' not in current_test:
                    current_test['Description'] = []
                current_test['Description'].append(line)
            
            # Add content to current section
            elif current_test:
                if current_section:
                    if current_section not in current_test:
                        current_test[current_section] = []
                    current_test[current_section].append(line)
                else:
                    if 'Description' not in current_test:
                        current_test['Description'] = []
                    current_test['Description'].append(line)
        
        if current_test:
            data.append(current_test)
        
        return data if data else [{'Test Case ID': 'Unable to parse'}]
    
    def parse_parent_id_assignment(self, assignment: str) -> Dict[str, str]:
        """Parse flexible Parent ID assignment"""
        mapping = {}
        
        # Case: "ADXP-5439 for all"
        if "for all" in assignment.lower():
            parent_id = assignment.split("for all")[0].strip()
            mapping['all'] = parent_id
            return mapping
        
        # Case: "ADXP-5439 for 001-005, ADXP-5440 for 006-010"
        try:
            parts = assignment.split(',')
            for part in parts:
                part = part.strip()
                if 'for' in part:
                    parent_id, range_part = part.split('for')
                    parent_id = parent_id.strip()
                    range_part = range_part.strip()
                    
                    if '-' in range_part:
                        start, end = range_part.split('-')
                        mapping[f"range:{start.strip()}:{end.strip()}"] = parent_id
        except:
            print(f"Warning: Could not parse parent ID assignment: {assignment}")
            mapping['all'] = assignment.strip()
        
        return mapping
    
    def get_parent_id(self, test_id: str, mapping: Dict[str, str]) -> str:
        """Get parent ID for a test case"""
        if not mapping:
            return ""
        
        if 'all' in mapping:
            return mapping['all']
        
        for key, parent_id in mapping.items():
            if key.startswith('range:'):
                _, start, end = key.split(':')
                if start <= test_id <= end:
                    return parent_id
        
        return ""
    
    def ask_for_input(self) -> tuple:
        """Interactively ask for required inputs"""
        print("\n" + "="*60)
        print("TEST SCENARIO CSV CONVERTER - INTERACTIVE MODE")
        print("="*60)
        
        # Parent ID
        print("\nParent ID Assignment Options:")
        print("  1. 'ADXP-5439 for all'")
        print("  2. 'ADXP-5439 for 001-005, ADXP-5440 for 006-010'")
        parent_ids = input("\nEnter Parent ID assignment: ").strip()
        
        # Assignee
        assignee = input("Enter Jira assignee (username): ").strip()
        
        # Test Type
        test_type = input("Enter test type (Manual/Automated): ").strip() or "Manual"
        
        # Fix Version
        fix_version = input("Enter Fix Version (or press Enter to skip): ").strip()
        
        return parent_ids, assignee, test_type, fix_version
    
    def generate_csv_rows(self, test_cases: List[Dict], 
                         parent_ids: Dict[str, str],
                         assignee: str, test_type: str, 
                         fix_version: str) -> List[Dict]:
        """Generate CSV rows"""
        rows = []
        
        for test_case in test_cases:
            row = {header: self.DEFAULTS.get(header, "") for header in self.HEADERS}
            
            # Extract test ID
            test_id = str(test_case.get('Test Case ID', '001')).strip()
            test_id = re.sub(r'\D', '', test_id)[:3].zfill(3)
            
            # Apply parent ID
            parent_id = self.get_parent_id(test_id, parent_ids)
            
            # Fill in row
            row['Parent ID'] = parent_id
            row['Assignee'] = assignee
            row['TestType'] = test_type
            row['Fix Version'] = fix_version
            row['Test Scenario/Test Case ID'] = test_id
            
            # Preserve test data
            for key in ['Pre-Condition', 'Test Steps', 'Expected Result', 'Test Data', 'Acceptance Criteria', 'User Story ID']:
                if key in test_case:
                    row[key] = str(test_case[key]).strip()
                elif key.replace('-', ' ').lower() in str(test_case).lower():
                    # Try to find similar keys
                    for k, v in test_case.items():
                        if key.lower() in k.lower():
                            row[key] = str(v).strip()
                            break
            
            # Description
            description = str(test_case.get('Description', 
                                           test_case.get('Test Scenario Description', ''))).strip()
            row['Test Scenario Description'] = description
            
            # Summary
            summary = description[:150] if description else f"Test Case {test_id}"
            row['Summary'] = f"Test Case {test_id} - {summary}" if summary else f"Test Case {test_id}"
            
            rows.append(row)
        
        return rows
    
    def write_csv(self, rows: List[Dict], output_file: str) -> None:
        """Write CSV file UTF-8"""
        try:
            with open(output_file, 'w', newline='', encoding='utf-8') as f:
                writer = csv.DictWriter(f, fieldnames=self.HEADERS, extrasaction='ignore')
                writer.writeheader()
                writer.writerows(rows)
            
            print(f"\n✓ SUCCESS: CSV generated")
            print(f"  File: {output_file}")
            print(f"  Test cases: {len(rows)}")
            print(f"  Format: UTF-8")
            print(f"  Ready for Jira import")
        except Exception as e:
            print(f"Error writing CSV: {e}")
            sys.exit(1)


def main():
    parser = argparse.ArgumentParser(
        description='Convert test scenarios to Jira CSV format',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''Examples:
  python test_scenario_converter.py --input scenarios.xlsx --parent-ids "ADXP-5439 for all" --assignee john_smith
  python test_scenario_converter.py --input data.csv --parent-ids "ADXP-5439 for 001-005, ADXP-5440 for 006-010"
  python test_scenario_converter.py --input scenarios.xlsx  # Interactive mode
        '''
    )
    
    parser.add_argument('--input', required=True, help='Input file or text')
    parser.add_argument('--output', default='test_scenarios.csv', help='Output file')
    parser.add_argument('--parent-ids', help='Parent ID assignment')
    parser.add_argument('--assignee', help='Jira assignee')
    parser.add_argument('--test-type', default='Manual', help='Test type')
    parser.add_argument('--fix-version', help='Fix version')
    
    args = parser.parse_args()
    
    converter = TestScenarioConverter()
    
    print("="*60)
    print("TEST SCENARIO CSV CONVERTER")
    print("="*60)
    
    # Read input
    print(f"\nReading: {args.input}")
    test_cases = converter.read_input(args.input)
    print(f"Found {len(test_cases)} test case(s)")
    
    # Get inputs
    if args.parent_ids:
        parent_ids = converter.parse_parent_id_assignment(args.parent_ids)
        assignee = args.assignee or "unassigned"
        test_type = args.test_type
        fix_version = args.fix_version or ""
    else:
        # Interactive
        parent_id_str, assignee, test_type, fix_version = converter.ask_for_input()
        parent_ids = converter.parse_parent_id_assignment(parent_id_str)
    
    # Generate and write
    rows = converter.generate_csv_rows(test_cases, parent_ids, assignee, test_type, fix_version)
    converter.write_csv(rows, args.output)
    print(f"\n✓ Next: Import {args.output} to Jira via Issues → Import")


if __name__ == '__main__':
    main()